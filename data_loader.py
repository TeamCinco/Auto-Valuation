"""
data_loader.py
Handles reading and parsing market cap and cash flow files (CSV or XLSX).
"""

import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime


def parse_mcap_value(val):
    """Parse market cap value: raw numbers (8698100000) or abbreviated (11.98B, 450.5M)."""
    s = str(val).strip().replace(',', '').replace('$', '')
    if not s or s.lower() == 'nan':
        return np.nan
    multipliers = {'T': 1e12, 'B': 1e9, 'M': 1e6, 'K': 1e3}
    suffix = s[-1].upper()
    if suffix in multipliers:
        try:
            return float(s[:-1]) * multipliers[suffix]
        except ValueError:
            return np.nan
    try:
        return float(s)
    except ValueError:
        return np.nan


def parse_cf_value(val):
    """Parse cash flow value: handles commas, parens for negatives, spaces."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if isinstance(val, float) and np.isnan(val):
            return None
        return float(val)
    s = str(val).strip()
    if not s or s.lower() == 'nan' or s == '-':
        return None
    negative = False
    if s.startswith('(') and s.endswith(')'):
        negative = True
        s = s[1:-1]
    s = s.replace(',', '').replace('$', '').replace(' ', '')
    try:
        result = float(s)
        return -result if negative else result
    except ValueError:
        return None


def quarter_end_date(dt):
    """Map a date to its quarter-end date."""
    m, y = dt.month, dt.year
    if m <= 3:
        return pd.Timestamp(y, 3, 31)
    elif m <= 6:
        return pd.Timestamp(y, 6, 30)
    elif m <= 9:
        return pd.Timestamp(y, 9, 30)
    else:
        return pd.Timestamp(y, 12, 31)


def load_market_cap(filepath):
    """Load daily market cap file -> quarterly averages in millions."""
    ext = os.path.splitext(filepath)[1].lower()
    df = pd.read_excel(filepath) if ext == '.xlsx' else pd.read_csv(filepath)

    df.columns = [c.strip() for c in df.columns]
    mcap_col = [c for c in df.columns if 'market' in c.lower() and 'cap' in c.lower()]
    if not mcap_col:
        mcap_col = [c for c in df.columns if c != 'Date' and 'change' not in c.lower()]
    mcap_col = mcap_col[0]
    date_col = [c for c in df.columns if 'date' in c.lower()][0]

    df[date_col] = pd.to_datetime(df[date_col], format='mixed')
    df[mcap_col] = df[mcap_col].apply(parse_mcap_value)
    df['mcap_mm'] = df[mcap_col] / 1_000_000
    df['quarter_end'] = df[date_col].apply(quarter_end_date)

    return df.groupby('quarter_end')['mcap_mm'].mean().sort_index()


def load_cash_flow(filepath):
    """Load cash flow file -> dict of {quarter_date: {net_income, fcf}}."""
    ext = os.path.splitext(filepath)[1].lower()

    if ext == '.xlsx':
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        data = {}
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            label = row[0]
            if label is not None:
                data[str(label).strip()] = list(row[1:])
    else:
        df = pd.read_csv(filepath, header=None)
        data = {}
        for _, row in df.iterrows():
            label = row.iloc[0]
            if pd.notna(label):
                data[str(label).strip()] = list(row.iloc[1:])

    raw_dates = data.get('Year Ending', [])
    dates = []
    for d in raw_dates:
        if d is None or (isinstance(d, float) and np.isnan(d)):
            dates.append(None)
        elif isinstance(d, datetime):
            dates.append(pd.Timestamp(d))
        elif isinstance(d, str):
            try:
                dates.append(pd.Timestamp(d))
            except Exception:
                dates.append(None)
        else:
            dates.append(None)

    net_income_raw = data.get('Net Income', [])
    fcf_raw = data.get('Free Cash Flow', [])

    result = {}
    for i, dt in enumerate(dates):
        if dt is None:
            continue
        ni_val = parse_cf_value(net_income_raw[i]) if i < len(net_income_raw) else None
        fcf_val = parse_cf_value(fcf_raw[i]) if i < len(fcf_raw) else None
        result[dt] = {'net_income': ni_val, 'fcf': fcf_val}

    return result