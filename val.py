import os
import sys
import pandas as pd
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from datetime import datetime
import warnings
warnings.filterwarnings('ignore')


def load_market_cap(filepath):
    """Load daily market cap (CSV or XLSX) -> quarterly averages in millions."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.xlsx':
        df = pd.read_excel(filepath)
    else:
        df = pd.read_csv(filepath)
    
    # Normalize column names
    df.columns = [c.strip() for c in df.columns]
    
    # Find market cap column (could be 'Market Cap' or similar)
    mcap_col = [c for c in df.columns if 'market' in c.lower() and 'cap' in c.lower()]
    if not mcap_col:
        mcap_col = [c for c in df.columns if c != 'Date' and 'change' not in c.lower()]
    mcap_col = mcap_col[0]
    
    date_col = [c for c in df.columns if 'date' in c.lower()][0]
    
    df[date_col] = pd.to_datetime(df[date_col], format='mixed')
    
    # Parse market cap — handles both raw numbers (8698100000) and abbreviated (11.98B, 450.5M, 1.2T)
    def parse_mcap(val):
        s = str(val).strip().replace(',', '').replace('$', '')
        if not s or s.lower() == 'nan':
            return np.nan
        multipliers = {'T': 1_000_000_000_000, 'B': 1_000_000_000, 'M': 1_000_000, 'K': 1_000}
        suffix = s[-1].upper()
        if suffix in multipliers:
            try:
                return float(s[:-1]) * multipliers[suffix]
            except ValueError:
                return np.nan
        try:
            return float(s)
        except ValueError:
            return np.nan
    
    df[mcap_col] = df[mcap_col].apply(parse_mcap)
    
    # Convert to millions
    df['mcap_mm'] = df[mcap_col] / 1_000_000
    
    # Assign quarter-end date for each row
    df['quarter_end'] = df[date_col].apply(quarter_end_date)
    
    # Group by quarter and average
    quarterly = df.groupby('quarter_end')['mcap_mm'].mean().sort_index()
    
    return quarterly


def quarter_end_date(dt):
    """Map a date to its quarter-end date."""
    month = dt.month
    year = dt.year
    if month <= 3:
        return pd.Timestamp(year, 3, 31)
    elif month <= 6:
        return pd.Timestamp(year, 6, 30)
    elif month <= 9:
        return pd.Timestamp(year, 9, 30)
    else:
        return pd.Timestamp(year, 12, 31)


def parse_cf_value(val):
    """Parse a cash flow value that may be a string like ' 1,252 ' or '(2,722)' or already numeric."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if np.isnan(val) if isinstance(val, float) else False:
            return None
        return float(val)
    s = str(val).strip()
    if not s or s.lower() == 'nan' or s == '-':
        return None
    # Handle parentheses for negatives: (2,722) -> -2722
    negative = False
    if s.startswith('(') and s.endswith(')'):
        negative = True
        s = s[1:-1]
    s = s.replace(',', '').replace('$', '').replace(' ', '')
    try:
        result = float(s)
        return -result if negative else result
    except ValueError:
        return None


def load_cash_flow(filepath):
    """Load cash flow (CSV or XLSX) -> dict with dates, net_income, fcf."""
    ext = os.path.splitext(filepath)[1].lower()
    
    if ext == '.xlsx':
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active
        data = {}
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            label = row[0]
            if label is not None:
                data[str(label).strip()] = list(row[1:])
    else:
        # CSV
        df = pd.read_csv(filepath, header=None)
        data = {}
        for _, row in df.iterrows():
            label = row.iloc[0]
            if pd.notna(label):
                data[str(label).strip()] = list(row.iloc[1:])
    
    # Parse dates from 'Year Ending' row
    raw_dates = data.get('Year Ending', [])
    dates = []
    for d in raw_dates:
        if d is None or (isinstance(d, float) and np.isnan(d)):
            dates.append(None)
        elif isinstance(d, datetime):
            dates.append(pd.Timestamp(d))
        elif isinstance(d, str):
            try:
                dates.append(pd.Timestamp(d))
            except:
                dates.append(None)
        else:
            dates.append(None)
    
    # Get Net Income and Free Cash Flow rows
    net_income_raw = data.get('Net Income', [])
    fcf_raw = data.get('Free Cash Flow', [])
    
    # Build series aligned to dates, parsing values to float
    result = {}
    for i, dt in enumerate(dates):
        if dt is None:
            continue
        ni_val = parse_cf_value(net_income_raw[i]) if i < len(net_income_raw) else None
        fcf_val = parse_cf_value(fcf_raw[i]) if i < len(fcf_raw) else None
        result[dt] = {
            'net_income': ni_val,
            'fcf': fcf_val
        }
    
    return result


def process_ticker(ticker, mcap_file, cf_file):
    """Process one ticker's files -> dict with quarterly data."""
    print(f"  Market cap:  {os.path.basename(mcap_file)}")
    quarterly_mcap = load_market_cap(mcap_file)
    
    print(f"  Cash flow:   {os.path.basename(cf_file)}")
    cf_data = load_cash_flow(cf_file)
    
    # Build unified quarterly timeline
    all_quarters = sorted(set(list(quarterly_mcap.index) + list(cf_data.keys())))
    
    result = {
        'ticker': ticker,
        'quarters': [],
        'net_income': [],
        'fcf': [],
        'avg_mcap': [],
        'avg_pe': [],
        'avg_pfcf': []
    }
    
    for q in all_quarters:
        ni = cf_data.get(q, {}).get('net_income')
        fcf = cf_data.get(q, {}).get('fcf')
        mcap = quarterly_mcap.get(q)
        
        # Only include quarters where we have at least CF data
        if ni is None and fcf is None:
            continue
            
        result['quarters'].append(q)
        result['net_income'].append(ni)
        result['fcf'].append(fcf)
        result['avg_mcap'].append(round(mcap, 0) if mcap is not None and not np.isnan(mcap) else None)
        
        # Compute PE and P/FCF
        pe = None
        pfcf = None
        if mcap is not None and not np.isnan(mcap):
            if ni is not None and ni != 0:
                pe = round(float(mcap) / float(ni), 2)
            if fcf is not None and fcf != 0:
                pfcf = round(float(mcap) / float(fcf), 2)
        
        result['avg_pe'].append(pe)
        result['avg_pfcf'].append(pfcf)
    
    return result


def format_quarter_label(dt):
    """Format quarter date for column header."""
    return dt.strftime('%Y-%m-%d')


def build_valuation_sheet(tickers_data, output_path):
    """Build the final valuation Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Valuation"
    
    # --- Styles ---
    header_font = Font(name='Arial', bold=True, size=10)
    ticker_font = Font(name='Arial', bold=True, size=10, color='000080')
    data_font = Font(name='Arial', size=9)
    num_fmt_int = '#,##0'
    num_fmt_dec = '#,##0.00'
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    ticker_fill = PatternFill('solid', fgColor='E2EFDA')
    thin_border = Border(
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    # ============================================================
    # SECTION 1: Summary table (P/FCF Valuation)
    # ============================================================
    ws.cell(row=1, column=1, value="P FCF Valuation").font = Font(name='Arial', bold=True, size=12)
    
    # Headers
    summary_headers = ['Ticker', 'Market Cap', 'TTM FCF', 'P/FCF', '5Y Avg', '10Y Avg']
    for j, h in enumerate(summary_headers):
        cell = ws.cell(row=2, column=j+1, value=h)
        cell.font = header_font
        cell.fill = header_fill
    
    summary_start_row = 3
    for i, td in enumerate(tickers_data):
        row = summary_start_row + i
        ticker = td['ticker']
        
        # Latest values
        latest_mcap = None
        latest_fcf = None
        latest_pfcf = None
        
        for idx in range(len(td['quarters'])-1, -1, -1):
            if td['avg_mcap'][idx] is not None:
                latest_mcap = td['avg_mcap'][idx]
                break
        for idx in range(len(td['quarters'])-1, -1, -1):
            if td['fcf'][idx] is not None:
                latest_fcf = td['fcf'][idx]
                break
        for idx in range(len(td['quarters'])-1, -1, -1):
            if td['avg_pfcf'][idx] is not None:
                latest_pfcf = td['avg_pfcf'][idx]
                break
        
        # 5Y and 10Y averages of P/FCF
        pfcf_vals = [v for v in td['avg_pfcf'] if v is not None]
        avg_5y = round(np.mean(pfcf_vals[-20:]), 2) if len(pfcf_vals) >= 4 else None  # ~20 quarters = 5 years
        avg_10y = round(np.mean(pfcf_vals[-40:]), 2) if len(pfcf_vals) >= 4 else None  # ~40 quarters = 10 years
        
        ws.cell(row=row, column=1, value=ticker).font = ticker_font
        c = ws.cell(row=row, column=2, value=latest_mcap)
        c.number_format = num_fmt_int
        c.font = data_font
        c = ws.cell(row=row, column=3, value=latest_fcf)
        c.number_format = num_fmt_int
        c.font = data_font
        c = ws.cell(row=row, column=4, value=latest_pfcf)
        c.number_format = num_fmt_dec
        c.font = data_font
        c = ws.cell(row=row, column=5, value=avg_5y)
        c.number_format = num_fmt_dec
        c.font = data_font
        c = ws.cell(row=row, column=6, value=avg_10y)
        c.number_format = num_fmt_dec
        c.font = data_font
    
    # ============================================================
    # SECTION 2: Per-ticker quarterly detail blocks
    # ============================================================
    current_row = summary_start_row + len(tickers_data) + 2
    
    for td in tickers_data:
        ticker = td['ticker']
        quarters = td['quarters']
        n_quarters = len(quarters)
        
        if n_quarters == 0:
            continue
        
        # Ticker header row with quarter dates
        ws.cell(row=current_row, column=1, value=ticker).font = ticker_font
        ws.cell(row=current_row, column=1).fill = ticker_fill
        for j, q in enumerate(quarters):
            cell = ws.cell(row=current_row, column=j+2, value=format_quarter_label(q))
            cell.font = Font(name='Arial', size=8, bold=True)
            cell.fill = ticker_fill
            cell.alignment = Alignment(horizontal='center')
        
        # Row labels and data
        row_labels = ['Net Income TTM', 'FCF TTM', 'Average Market Cap', 'Average PE', 'Average P/FCF']
        row_data = [td['net_income'], td['fcf'], td['avg_mcap'], td['avg_pe'], td['avg_pfcf']]
        
        for k, (label, values) in enumerate(zip(row_labels, row_data)):
            r = current_row + 1 + k
            ws.cell(row=r, column=1, value=label).font = data_font
            
            for j, val in enumerate(values):
                cell = ws.cell(row=r, column=j+2, value=val)
                cell.font = data_font
                if label in ('Net Income TTM', 'FCF TTM', 'Average Market Cap'):
                    cell.number_format = num_fmt_int
                else:
                    cell.number_format = num_fmt_dec
                cell.alignment = Alignment(horizontal='right')
                cell.border = thin_border
        
        current_row += 1 + len(row_labels) + 1  # +1 for blank separator row
    
    # Column widths
    ws.column_dimensions['A'].width = 20
    for col in range(2, 50):
        ws.column_dimensions[get_column_letter(col)].width = 14
    
    wb.save(output_path)
    print(f"\nOutput saved to: {output_path}")


def discover_tickers(input_dir):
    """Scan a flat folder for TICKER-market-cap.* and TICKER-cash-flow-statement-ttm.* pairs.
    
    Naming convention:
      TICKER-market-cap.csv  or  TICKER-market-cap.xlsx
      TICKER-cash-flow-statement-ttm.csv  or  TICKER-cash-flow-statement-ttm.xlsx
    """
    files = [f for f in os.listdir(input_dir) if not f.startswith('.')]
    
    mcap_files = {}   # ticker -> filepath
    cf_files = {}     # ticker -> filepath
    
    for f in files:
        fl = f.lower()
        if not (fl.endswith('.csv') or fl.endswith('.xlsx')):
            continue
        
        # Match: TICKER-market-cap.ext
        if '-market-cap.' in fl:
            ticker = f[:f.lower().index('-market-cap.')].upper()
            mcap_files[ticker] = os.path.join(input_dir, f)
        
        # Match: TICKER-cash-flow-statement-ttm.ext
        elif '-cash-flow-statement-ttm.' in fl:
            ticker = f[:f.lower().index('-cash-flow-statement-ttm.')].upper()
            cf_files[ticker] = os.path.join(input_dir, f)
    
    # Find tickers that have BOTH files
    all_tickers = sorted(set(mcap_files.keys()) & set(cf_files.keys()))
    
    # Report missing pairs
    mcap_only = set(mcap_files.keys()) - set(cf_files.keys())
    cf_only = set(cf_files.keys()) - set(mcap_files.keys())
    if mcap_only:
        print(f"WARNING: Market cap found but no cash flow for: {', '.join(sorted(mcap_only))}")
    if cf_only:
        print(f"WARNING: Cash flow found but no market cap for: {', '.join(sorted(cf_only))}")
    
    return [(t, mcap_files[t], cf_files[t]) for t in all_tickers]


def main():
    input_dir = sys.argv[1] if len(sys.argv) > 1 else 'input_data'
    output_file = sys.argv[2] if len(sys.argv) > 2 else 'valuation_output.xlsx'
    
    if not os.path.isdir(input_dir):
        print(f"ERROR: Input directory '{input_dir}' not found.")
        print(f"\nCreate the folder and drop your files in:")
        print(f"  {input_dir}/")
        print(f"    SNAP-market-cap.csv")
        print(f"    SNAP-cash-flow-statement-ttm.xlsx")
        print(f"    META-market-cap.csv")
        print(f"    META-cash-flow-statement-ttm.csv")
        print(f"    GOOG-market-cap.xlsx")
        print(f"    GOOG-cash-flow-statement-ttm.xlsx")
        print(f"    ...")
        print(f"\nFiles can be .csv or .xlsx in any mix.")
        sys.exit(1)
    
    pairs = discover_tickers(input_dir)
    
    if not pairs:
        print(f"ERROR: No valid ticker pairs found in '{input_dir}/'")
        print(f"\nExpected naming convention:")
        print(f"  TICKER-market-cap.csv  (or .xlsx)")
        print(f"  TICKER-cash-flow-statement-ttm.csv  (or .xlsx)")
        print(f"\nExample: SNAP-market-cap.csv + SNAP-cash-flow-statement-ttm.xlsx")
        sys.exit(1)
    
    tickers = [t for t, _, _ in pairs]
    print(f"Found {len(pairs)} tickers: {', '.join(tickers)}\n")
    
    all_data = []
    for ticker, mcap_path, cf_path in pairs:
        print(f"Processing {ticker}...")
        result = process_ticker(ticker, mcap_path, cf_path)
        if result:
            all_data.append(result)
    
    if not all_data:
        print("ERROR: No valid data processed.")
        sys.exit(1)
    
    build_valuation_sheet(all_data, output_file)
    print("Done!")


if __name__ == '__main__':
    main()